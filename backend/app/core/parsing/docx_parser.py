"""
DOCX text extraction.

Responsibility is deliberately narrow: turn a .docx file's bytes into
plain text, preserving paragraph and table-cell boundaries as newlines
so downstream section detection has clean line breaks to work with.
No structural interpretation happens here - that's resume_parser.py's
job.
"""

import io

from docx import Document


class DocxParseError(ValueError):
    """Raised when a .docx file cannot be read/parsed."""


def extract_text(content: bytes) -> str:
    """Extract plain text from DOCX file bytes.

    Raises DocxParseError if the file is not a valid, readable .docx.
    """
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:  # python-docx raises varied exception types
        raise DocxParseError(f"Could not read DOCX file: {exc}") from exc

    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(text)

    return "\n".join(lines)
